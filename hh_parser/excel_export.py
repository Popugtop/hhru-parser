"""Excel-экспорт (раздел 2а задания): лист «Выборка» с подсветкой расхождений
и выпадающим списком, плюс листы «Количества по сферам», «Сводка», «Отсев»,
«Параметры сбора», «Словарь колонок»."""
from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from hh_parser.aggregate import SphereCounts
from hh_parser.codes import EDUCATION_FILTER
from hh_parser.columns import COLUMNS, ColumnSpec
from hh_parser.format_values import display_value
from hh_parser.rows import FinalRow
from hh_parser.sample import SampleResult
from hh_parser.summary import compute_group_summaries

HEADER_FONT = Font(bold=True)
MISMATCH_FILL = PatternFill(start_color="FFFCE8B2", end_color="FFFCE8B2", fill_type="solid")
HYPERLINK_FONT = Font(color="0563C1", underline="single")
NO_WRAP = Alignment(wrap_text=False)

EDUCATION_GROUP_ORDER = ["higher", "not_required_or_not_specified", "special_secondary"]
DESCRIPTION_COLUMN_WIDTH = 50
MAX_COLUMN_WIDTH = 60
PERCENT_FORMAT = '0.0"%"'
MONEY_FORMAT = "#,##0"
INT_FORMAT = "0"
DATE_FORMAT = "DD.MM.YYYY"


def sort_key_v1(row: FinalRow) -> tuple[int, str, float]:
    order = (
        EDUCATION_GROUP_ORDER.index(row.education_group_filter)
        if row.education_group_filter in EDUCATION_GROUP_ORDER
        else len(EDUCATION_GROUP_ORDER)
    )
    salary = row.salary_net_estimate if row.salary_net_estimate is not None else -1
    return (order, row.sphere or "", salary)


def excel_value_and_format(spec: ColumnSpec, row: FinalRow) -> tuple[Any, str | None]:
    if spec.kind == "money":
        raw = spec.getter(row)
        return (raw if raw is not None else None), MONEY_FORMAT
    if spec.kind == "int":
        raw = spec.getter(row)
        return (raw if raw is not None else None), INT_FORMAT
    if spec.kind == "date":
        raw = spec.getter(row)
        if not raw:
            return None, None
        year, month, day = (int(p) for p in raw.split("-"))
        return datetime.date(year, month, day), DATE_FORMAT
    return display_value(spec, row), None


def write_header(ws: Worksheet, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def autosize_columns(
    ws: Worksheet,
    headers: list[str],
    rows_values: list[list[Any]],
    narrow_columns: dict[int, int] | None = None,
) -> None:
    narrow_columns = narrow_columns or {}
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        if col_idx in narrow_columns:
            ws.column_dimensions[letter].width = narrow_columns[col_idx]
            continue
        max_len = len(str(header))
        for row_values in rows_values:
            value = row_values[col_idx - 1]
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max_len + 2, MAX_COLUMN_WIDTH)


def _write_vacancies_sheet(wb: Workbook, rows: list[FinalRow]) -> Worksheet:
    ws = wb.active
    ws.title = "Выборка"
    headers = [spec.header for spec in COLUMNS]
    write_header(ws, headers)

    sorted_rows = sorted(rows, key=sort_key_v1)
    link_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "link")
    verified_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "verified_manually")
    description_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "description")

    all_row_values: list[list[Any]] = []
    for row_idx, row in enumerate(sorted_rows, start=2):
        row_values: list[Any] = []
        for col_idx, spec in enumerate(COLUMNS, start=1):
            value, number_format = excel_value_and_format(spec, row)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if number_format:
                cell.number_format = number_format
            if col_idx == link_col_idx and value:
                cell.hyperlink = value
                cell.font = HYPERLINK_FONT
            if col_idx == description_col_idx:
                cell.alignment = NO_WRAP
            row_values.append(value)
        if row.education_mismatch:
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = MISMATCH_FILL
        all_row_values.append(row_values)

    narrow = {description_col_idx: DESCRIPTION_COLUMN_WIDTH}
    autosize_columns(ws, headers, all_row_values, narrow_columns=narrow)

    dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
    ws.add_data_validation(dv)
    last_row = max(len(sorted_rows) + 1, 2)
    verified_letter = get_column_letter(verified_col_idx)
    dv.add(f"{verified_letter}2:{verified_letter}{last_row}")

    return ws


def write_counts_sheet(wb: Workbook, sphere_rows: list[SphereCounts]) -> Worksheet:
    ws = wb.create_sheet("Количества по сферам")
    headers = [
        "Сфера", "Всего вакансий", "Высшее", "СПО",
        "Не требуется или не указано", 'Доля «Высшее», %',
    ]
    write_header(ws, headers)
    all_row_values = []
    for row_idx, row in enumerate(sphere_rows, start=2):
        values = [
            row.label, row.total_found, row.higher,
            row.special_secondary, row.not_required_or_not_specified, row.higher_share_pct,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (2, 3, 4, 5):
                cell.number_format = MONEY_FORMAT
            elif col_idx == 6 and value is not None:
                cell.number_format = PERCENT_FORMAT
        all_row_values.append(values)
    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_summary_sheet(wb: Workbook, rows: list[FinalRow]) -> Worksheet:
    ws = wb.create_sheet("Сводка")
    headers = [
        "Группа образования", "Число вакансий", "Медиана зарплаты на руки, руб.",
        "Средняя зарплата на руки, руб.", "Доля с ТК РФ, %", "Доля с ДМС, %",
        "Доля с обучением/наставничеством, %", "Доля с удалёнкой, %", "Доля бюджетного сектора, %",
    ]
    write_header(ws, headers)
    summaries = compute_group_summaries(rows)
    all_row_values = []
    for row_idx, s in enumerate(summaries, start=2):
        values = [
            s.group_label, s.count, s.median_salary_net, s.mean_salary_net,
            s.share_labor_contract, s.share_dms, s.share_training, s.share_remote, s.share_budget,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (3, 4) and value is not None:
                cell.number_format = MONEY_FORMAT
            elif col_idx in (5, 6, 7, 8, 9) and value is not None:
                cell.number_format = PERCENT_FORMAT
        all_row_values.append(values)
    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_exclusions_sheet(wb: Workbook, sample_results: dict[str, SampleResult]) -> Worksheet:
    ws = wb.create_sheet("Отсев")
    headers = ["Группа образования", "Размер пула", "Запрошено", "Собрано", "Причина отсева", "Количество"]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    row_idx = 2
    for group_filter, result in sample_results.items():
        label = EDUCATION_FILTER.get(group_filter, group_filter)
        base = [label, result.pool_size, result.requested_quota, len(result.collected)]
        reasons = result.excluded_counts
        if not reasons:
            values = base + ["", ""]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            all_row_values.append(values)
            row_idx += 1
            continue
        for reason, count in reasons.items():
            values = base + [reason, count]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            all_row_values.append(values)
            row_idx += 1
    autosize_columns(ws, headers, all_row_values)
    return ws


def write_params_sheet(wb: Workbook, run_info: dict[str, Any]) -> Worksheet:
    ws = wb.create_sheet("Параметры сбора")
    headers = ["Параметр", "Значение"]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    for row_idx, (key, value) in enumerate(run_info.items(), start=2):
        if isinstance(value, dict):
            display: Any = ", ".join(f"{k}={v}" for k, v in value.items())
        elif isinstance(value, (list, tuple)):
            display = ", ".join(str(v) for v in value)
        else:
            display = value
        ws.cell(row=row_idx, column=1, value=key)
        ws.cell(row=row_idx, column=2, value=display)
        all_row_values.append([key, display])
    autosize_columns(ws, headers, all_row_values, narrow_columns={2: 70})
    return ws


def _write_dictionary_sheet(wb: Workbook) -> Worksheet:
    ws = wb.create_sheet("Словарь колонок")
    headers = ["Колонка", "Что в ней и откуда взято"]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    for row_idx, spec in enumerate(COLUMNS, start=2):
        ws.cell(row=row_idx, column=1, value=spec.header)
        ws.cell(row=row_idx, column=2, value=spec.source)
        all_row_values.append([spec.header, spec.source])
    autosize_columns(ws, headers, all_row_values, narrow_columns={2: 80})
    return ws


def export_counts_workbook(sphere_rows: list[SphereCounts], output_path: str | Path) -> None:
    """Отдельная книга с одним листом «Количества по сферам» (совместимость
    с именованием output/counts_by_sphere.xlsx из раздела 2 задания)."""
    wb = Workbook()
    wb.remove(wb.active)
    write_counts_sheet(wb, sphere_rows)
    wb.save(output_path)


def export_workbook(
    rows: list[FinalRow],
    sphere_rows: list[SphereCounts],
    sample_results: dict[str, SampleResult],
    run_info: dict[str, Any],
    output_path: str | Path,
) -> None:
    wb = Workbook()
    _write_vacancies_sheet(wb, rows)
    write_counts_sheet(wb, sphere_rows)
    _write_summary_sheet(wb, rows)
    _write_exclusions_sheet(wb, sample_results)
    write_params_sheet(wb, run_info)
    _write_dictionary_sheet(wb)
    wb.save(output_path)

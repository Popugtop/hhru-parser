"""Excel-экспорт v2 (изменение дизайна, запрос пользователя от 2026-10-07):
2 группы по образованию, сферы по белому списку, подсветка расхождений
образования (жёлтым) и C-level ролей (оранжевым) отдельно.

Переиспользует общую инфраструктуру из excel_export.py (шапка, автоширина,
форматы чисел/дат, лист «Количества по сферам», лист «Параметры сбора») —
только листы «Выборка», «Сводка», «Отсев», «Словарь колонок» специфичны
для v2-схемы данных.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from hh_parser.aggregate import SphereCounts
from hh_parser.columns_v2 import COLUMNS
from hh_parser.excel_export import (
    DESCRIPTION_COLUMN_WIDTH,
    HYPERLINK_FONT,
    MISMATCH_FILL,
    MONEY_FORMAT,
    NO_WRAP,
    PERCENT_FORMAT,
    autosize_columns,
    excel_value_and_format,
    write_counts_sheet,
    write_header,
    write_params_sheet,
)
from hh_parser.rows_v2 import FinalRowV2
from hh_parser.sample_v2 import MatchedSampleResult
from hh_parser.summary_v2 import compute_group_summaries_v2

SENIOR_ROLE_FILL = PatternFill(start_color="FFFFD9B3", end_color="FFFFD9B3", fill_type="solid")

GROUP_2_ORDER = ["ВО требуется", "ВО не требуется"]


def sort_key_v2(row: FinalRowV2) -> tuple[int, str, float]:
    order = GROUP_2_ORDER.index(row.education_group_2) if row.education_group_2 in GROUP_2_ORDER else 99
    salary = row.salary_net_estimate if row.salary_net_estimate is not None else -1
    return (order, row.sphere or "", salary)


def _write_vacancies_sheet_v2(wb: Workbook, rows: list[FinalRowV2]) -> Worksheet:
    ws = wb.active
    ws.title = "Выборка"
    headers = [spec.header for spec in COLUMNS]
    write_header(ws, headers)

    sorted_rows = sorted(rows, key=sort_key_v2)
    link_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "link")
    verified_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "verified_manually")
    description_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "description")

    all_row_values: list[list[Any]] = []
    for row_idx, row in enumerate(sorted_rows, start=2):
        row_values: list[Any] = []
        for col_idx, spec in enumerate(COLUMNS, start=1):
            value, number_format = excel_value_and_format(spec, row)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if number_format:
                cell.number_format = number_format
            if col_idx == link_col_idx and value:
                cell.hyperlink = value
                cell.font = HYPERLINK_FONT
            if col_idx == description_col_idx:
                cell.alignment = NO_WRAP
            row_values.append(value)

        fill = MISMATCH_FILL if row.education_mismatch else (SENIOR_ROLE_FILL if row.is_senior_role else None)
        if fill is not None:
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill
        all_row_values.append(row_values)

    narrow = {description_col_idx: DESCRIPTION_COLUMN_WIDTH}
    autosize_columns(ws, headers, all_row_values, narrow_columns=narrow)

    dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
    ws.add_data_validation(dv)
    last_row = max(len(sorted_rows) + 1, 2)
    verified_letter = get_column_letter(verified_col_idx)
    dv.add(f"{verified_letter}2:{verified_letter}{last_row}")

    return ws


def _write_summary_sheet_v2(wb: Workbook, rows: list[FinalRowV2]) -> Worksheet:
    ws = wb.create_sheet("Сводка")
    headers = [
        "Группа ВО", "Число вакансий", "Медиана зарплаты на руки, руб.",
        "Средняя зарплата на руки, руб.", "Доля с ТК РФ, %", "Доля с ДМС, %",
        "Доля с обучением/наставничеством, %", "Доля с удалёнкой, %", "Доля бюджетного сектора, %",
    ]
    write_header(ws, headers)
    summaries = compute_group_summaries_v2(rows)
    all_row_values = []
    for row_idx, s in enumerate(summaries, start=2):
        values = [
            s.group_label, s.count, s.median_salary_net, s.mean_salary_net,
            s.share_labor_contract, s.share_dms, s.share_training, s.share_remote, s.share_budget,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (3, 4) and value is not None:
                cell.number_format = MONEY_FORMAT
            elif col_idx in (5, 6, 7, 8, 9) and value is not None:
                cell.number_format = PERCENT_FORMAT
        all_row_values.append(values)
    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_exclusions_sheet_v2(
    wb: Workbook, sample_results: dict[str, MatchedSampleResult]
) -> Worksheet:
    ws = wb.create_sheet("Отсев")
    headers = [
        "Группа ВО", "Сфера", "Размер пула (сферы)", "Цель (сферы)",
        "Собрано (сферы)", "Нехватка", "Причина отсева (по группе)", "Количество",
    ]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    row_idx = 2

    for group_label, result in sample_results.items():
        sphere_counts = result.sphere_counts()
        spheres = sorted(set(result.target_by_sphere) | set(result.pool_size_by_sphere) | set(sphere_counts))
        if not spheres:
            spheres = [""]
        for sphere in spheres:
            values = [
                group_label,
                sphere,
                result.pool_size_by_sphere.get(sphere, ""),
                result.target_by_sphere.get(sphere, ""),
                sphere_counts.get(sphere, 0),
                result.shortfall_by_sphere.get(sphere, ""),
                "",
                "",
            ]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            all_row_values.append(values)
            row_idx += 1

        reasons = result.excluded_counts
        if not reasons:
            continue
        for reason, count in reasons.items():
            values = [group_label, "", "", "", "", "", reason, count]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            all_row_values.append(values)
            row_idx += 1

    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_dictionary_sheet_v2(wb: Workbook) -> Worksheet:
    ws = wb.create_sheet("Словарь колонок")
    headers = ["Колонка", "Что в ней и откуда взято"]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    for row_idx, spec in enumerate(COLUMNS, start=2):
        ws.cell(row=row_idx, column=1, value=spec.header)
        ws.cell(row=row_idx, column=2, value=spec.source)
        all_row_values.append([spec.header, spec.source])
    autosize_columns(ws, headers, all_row_values, narrow_columns={2: 80})
    return ws


def export_workbook_v2(
    rows: list[FinalRowV2],
    sphere_rows: list[SphereCounts],
    sample_results: dict[str, MatchedSampleResult],
    run_info: dict[str, Any],
    output_path: str | Path,
) -> None:
    wb = Workbook()
    _write_vacancies_sheet_v2(wb, rows)
    write_counts_sheet(wb, sphere_rows)
    _write_summary_sheet_v2(wb, rows)
    _write_exclusions_sheet_v2(wb, sample_results)
    write_params_sheet(wb, run_info)
    _write_dictionary_sheet_v2(wb)
    wb.save(output_path)

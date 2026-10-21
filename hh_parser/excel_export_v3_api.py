"""Excel-экспорт v3 через официальный API (запрос пользователя 2026-10-09):
те же лист/колонки/подсветка, что в excel_export_v3.py, плюс лист
«Самопроверка» (часть А через API сравнена с website-версией v2, ±10%,
раздел 3 задания)."""
from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from hh_parser.aggregate import SphereCounts
from hh_parser.columns_v3_api import COLUMNS
from hh_parser.excel_export import (
    DESCRIPTION_COLUMN_WIDTH,
    HYPERLINK_FONT,
    MISMATCH_FILL,
    MONEY_FORMAT,
    NO_WRAP,
    PERCENT_FORMAT,
    autosize_columns,
    excel_value_and_format,
    write_counts_sheet,
    write_header,
    write_params_sheet,
)
from hh_parser.excel_export_v2 import SENIOR_ROLE_FILL, sort_key_v2
from hh_parser.rows_v2 import EDUCATION_GROUP_2_NOT_REQUIRED, EDUCATION_GROUP_2_REQUIRED
from hh_parser.rows_v3_api import FinalRowV3Api, assign_employer_sequence_numbers
from hh_parser.whitelist import ALL_SPHERES

GROUP_2_ORDER = [EDUCATION_GROUP_2_REQUIRED, EDUCATION_GROUP_2_NOT_REQUIRED]


def prepare_rows_for_export(rows: list[FinalRowV3Api]) -> list[FinalRowV3Api]:
    """Сортирует в финальном отображаемом порядке и проставляет
    employer_sequence_number — вызвать РОВНО ОДИН РАЗ перед экспортом в Excel
    и CSV, иначе номера работодателя в CSV останутся нулями."""
    sorted_rows = sorted(rows, key=sort_key_v2)
    assign_employer_sequence_numbers(sorted_rows)
    return sorted_rows


def _write_vacancies_sheet(wb: Workbook, rows: list[FinalRowV3Api]) -> Worksheet:
    ws = wb.active
    ws.title = "Выборка"
    headers = [spec.header for spec in COLUMNS]
    write_header(ws, headers)

    link_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "link")
    verified_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "verified_manually")
    description_col_idx = next(i for i, spec in enumerate(COLUMNS, start=1) if spec.key == "description")

    all_row_values: list[list[Any]] = []
    for row_idx, row in enumerate(rows, start=2):
        row_values: list[Any] = []
        for col_idx, spec in enumerate(COLUMNS, start=1):
            value, number_format = excel_value_and_format(spec, row)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if number_format:
                cell.number_format = number_format
            if col_idx == link_col_idx and value:
                cell.hyperlink = value
                cell.font = HYPERLINK_FONT
            if col_idx == description_col_idx:
                cell.alignment = NO_WRAP
            row_values.append(value)

        fill = MISMATCH_FILL if row.education_mismatch else (SENIOR_ROLE_FILL if row.is_senior_role else None)
        if fill is not None:
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill
        all_row_values.append(row_values)

    narrow = {description_col_idx: DESCRIPTION_COLUMN_WIDTH}
    autosize_columns(ws, headers, all_row_values, narrow_columns=narrow)

    dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
    ws.add_data_validation(dv)
    last_row = max(len(rows) + 1, 2)
    verified_letter = get_column_letter(verified_col_idx)
    dv.add(f"{verified_letter}2:{verified_letter}{last_row}")

    return ws


def _write_sphere_group_counts_sheet(wb: Workbook, rows: list[FinalRowV3Api]) -> Worksheet:
    ws = wb.create_sheet("Сферы x Группы ВО")
    headers = ["Сфера", EDUCATION_GROUP_2_REQUIRED, EDUCATION_GROUP_2_NOT_REQUIRED, "Итого"]
    write_header(ws, headers)

    counts: Counter[tuple[str, str]] = Counter()
    for row in rows:
        counts[(row.sphere or "[не определена]", row.education_group_2)] += 1

    spheres = list(ALL_SPHERES) + sorted({s for s, _ in counts if s not in ALL_SPHERES})
    all_row_values: list[list[Any]] = []
    row_idx = 2
    total_required = total_not_required = 0
    for sphere in spheres:
        required = counts.get((sphere, EDUCATION_GROUP_2_REQUIRED), 0)
        not_required = counts.get((sphere, EDUCATION_GROUP_2_NOT_REQUIRED), 0)
        if required == 0 and not_required == 0:
            continue
        total_required += required
        total_not_required += not_required
        values = [sphere, required, not_required, required + not_required]
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        all_row_values.append(values)
        row_idx += 1

    total_values = ["ИТОГО", total_required, total_not_required, total_required + total_not_required]
    for col_idx, value in enumerate(total_values, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)
    all_row_values.append(total_values)

    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_exclusions_sheet(
    wb: Workbook,
    excluded_counts: dict[str, int],
    pool_size_by_raw_filter: dict[str, int],
    required_excluded_counts: dict[str, int],
    not_required_excluded_counts: dict[str, int],
) -> Worksheet:
    ws = wb.create_sheet("Отсев")
    headers = ["Раздел", "Значение", "Количество"]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    row_idx = 2

    def add(section: str, label: str, count: int) -> None:
        nonlocal row_idx
        values = [section, label, count]
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        all_row_values.append(values)
        row_idx += 1

    for raw_filter, size in pool_size_by_raw_filter.items():
        add("Размер пула (до отбора)", raw_filter, size)
    for reason, count in excluded_counts.items():
        add("Причина отсева (сплошной сбор)", reason, count)
    for reason, count in required_excluded_counts.items():
        add("Причина отсева (подвыборка, «ВО требуется»)", reason, count)
    for reason, count in not_required_excluded_counts.items():
        add("Причина отсева (подвыборка, «ВО не требуется»)", reason, count)

    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_self_check_sheet(wb: Workbook, comparison_rows: list[dict[str, Any]]) -> Worksheet:
    ws = wb.create_sheet("Самопроверка")
    headers = [
        "Сфера", "Всего, API (с фильтром ЗП)", "Всего, сайт v2 (с фильтром ЗП)",
        "Расхождение, %", "Превышен порог ±10%",
    ]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    for row_idx, r in enumerate(comparison_rows, start=2):
        values = [
            r["sphere"], r["api_total"], r["web_total"], r["diff_pct"],
            "да" if r["exceeds"] else "нет",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (2, 3):
                cell.number_format = MONEY_FORMAT
            elif col_idx == 4 and value is not None:
                cell.number_format = PERCENT_FORMAT
        all_row_values.append(values)
    autosize_columns(ws, headers, all_row_values)
    return ws


def _write_dictionary_sheet(wb: Workbook) -> Worksheet:
    ws = wb.create_sheet("Словарь колонок")
    headers = ["Колонка", "Что в ней и откуда взято"]
    write_header(ws, headers)
    all_row_values: list[list[Any]] = []
    for row_idx, spec in enumerate(COLUMNS, start=2):
        ws.cell(row=row_idx, column=1, value=spec.header)
        ws.cell(row=row_idx, column=2, value=spec.source)
        all_row_values.append([spec.header, spec.source])
    autosize_columns(ws, headers, all_row_values, narrow_columns={2: 80})
    return ws


def export_workbook_v3_api(
    rows: list[FinalRowV3Api],
    sphere_rows: list[SphereCounts],
    excluded_counts: dict[str, int],
    pool_size_by_raw_filter: dict[str, int],
    required_excluded_counts: dict[str, int],
    not_required_excluded_counts: dict[str, int],
    self_check_rows: list[dict[str, Any]],
    run_info: dict[str, Any],
    output_path: str | Path,
) -> None:
    rows = prepare_rows_for_export(rows)
    wb = Workbook()
    _write_vacancies_sheet(wb, rows)
    write_counts_sheet(wb, sphere_rows)
    _write_sphere_group_counts_sheet(wb, rows)
    _write_exclusions_sheet(wb, excluded_counts, pool_size_by_raw_filter, required_excluded_counts, not_required_excluded_counts)
    _write_self_check_sheet(wb, self_check_rows)
    write_params_sheet(wb, run_info)
    _write_dictionary_sheet(wb)
    wb.save(output_path)

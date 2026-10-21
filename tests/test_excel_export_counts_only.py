import openpyxl

from hh_parser.aggregate import SphereCounts
from hh_parser.excel_export import export_counts_workbook


def test_export_counts_workbook_has_single_sheet(tmp_path):
    path = tmp_path / "counts.xlsx"
    rows = [SphereCounts("Все вакансии", 2183, 182, 106, 1944)]
    export_counts_workbook(rows, path)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Количества по сферам"]
    ws = wb["Количества по сферам"]
    assert ws.cell(row=2, column=1).value == "Все вакансии"
    assert ws.cell(row=2, column=2).value == 2183

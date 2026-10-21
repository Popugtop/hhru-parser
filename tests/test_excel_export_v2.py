from collections import Counter

import openpyxl
import pytest

from hh_parser.aggregate import SphereCounts
from hh_parser.excel_export_v2 import export_workbook_v2
from hh_parser.rows_v2 import EDUCATION_GROUP_2_NOT_REQUIRED, EDUCATION_GROUP_2_REQUIRED, FinalRowV2
from hh_parser.sample_v2 import GROUP_VO_NOT_REQUIRED, GROUP_VO_REQUIRED, MatchedSampleResult


def make_row(**overrides) -> FinalRowV2:
    base = dict(
        education_group_2=EDUCATION_GROUP_2_REQUIRED,
        education_group_filter="higher",
        education_group_filter_label="Высшее",
        vacancy_id=1,
        link="https://hh.ru/vacancy/1",
        name="Программист",
        employer="ООО Компания",
        employer_id=10,
        sector="коммерческий",
        profession_role="Программист, разработчик",
        role_category="Информационные технологии",
        sphere="IT",
        it_accreditation=True,
        publication_date="2026-10-01",
        salary_from=80000,
        salary_to=120000,
        gross=False,
        salary_net_estimate=100000,
        labor_contract=True,
        civil_law_contract_options=[],
        internship=False,
        work_formats=["HYBRID"],
        schedule=["FIVE_ON_TWO_OFF"],
        hours=["HOURS_8"],
        dms=True,
        training_mentorship=False,
        education_in_text="высшее",
        education_mismatch=False,
        is_senior_role=False,
    )
    base.update(overrides)
    return FinalRowV2(**base)


@pytest.fixture
def workbook_path(tmp_path):
    rows = [
        make_row(vacancy_id=1, sphere="IT", salary_net_estimate=100000),
        make_row(
            vacancy_id=2, education_group_2=EDUCATION_GROUP_2_NOT_REQUIRED,
            education_group_filter="special_secondary", education_group_filter_label="СПО",
            sphere="Право", salary_net_estimate=50000, education_mismatch=True,
            education_in_text="высшее",
        ),
        make_row(vacancy_id=3, sphere="Инженеры", salary_net_estimate=70000, is_senior_role=True),
    ]
    sphere_rows = [SphereCounts("IT", 50, 20, 10, 20)]
    sample_results = {
        GROUP_VO_REQUIRED: MatchedSampleResult(
            group_label=GROUP_VO_REQUIRED,
            sphere_by_vacancy_id={1: "IT", 3: "Инженеры"},
            pool_size_by_sphere={"IT": 30, "Инженеры": 10},
            excluded_counts=Counter({"Больше 3 вакансий этого работодателя в группе": 2}),
        ),
        GROUP_VO_NOT_REQUIRED: MatchedSampleResult(
            group_label=GROUP_VO_NOT_REQUIRED,
            sphere_by_vacancy_id={2: "Право"},
            target_by_sphere={"Право": 2, "IT": 1},
            pool_size_by_sphere={"Право": 5, "IT": 0},
            shortfall_by_sphere={"Право": 1, "IT": 1},
        ),
    }
    run_info = {"seed": 42, "area": 95}

    path = tmp_path / "vacancies_v2_test.xlsx"
    export_workbook_v2(rows, sphere_rows, sample_results, run_info, path)
    return path


def test_all_sheets_present(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    assert wb.sheetnames == [
        "Выборка", "Количества по сферам", "Сводка", "Отсев", "Параметры сбора", "Словарь колонок",
    ]


def test_vacancies_sheet_sorted_by_group_then_sphere_then_salary(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Выборка"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    id_col = headers.index("ID вакансии") + 1
    ids = [ws.cell(row=r, column=id_col).value for r in range(2, ws.max_row + 1)]
    # ВО требуется (1, 3) идёт раньше ВО не требуется (2); внутри группы —
    # по сфере (латиница "IT" < кириллица "Инженеры" по кодам Unicode)
    assert ids == [1, 3, 2]


def test_mismatch_and_senior_rows_get_distinct_highlight_colors(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Выборка"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    id_col = headers.index("ID вакансии") + 1

    fills_by_id = {}
    for row in ws.iter_rows(min_row=2):
        vid = row[id_col - 1].value
        fill = row[0].fill
        fills_by_id[vid] = fill.fgColor.rgb if fill and fill.fgColor else None

    assert fills_by_id[2] == "FFFCE8B2"  # расхождение образования — жёлтый
    assert fills_by_id[3] == "FFFFD9B3"  # C-level роль — оранжевый
    assert fills_by_id[1] in (None, "00000000")  # обычная строка — без подсветки


def test_no_bool_or_none_leak_in_vacancies_sheet(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Выборка"]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.value is not True
            assert cell.value is not False
            assert str(cell.value) != "nan"


def test_summary_sheet_has_two_groups(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Сводка"]
    labels = [ws.cell(row=r, column=1).value for r in range(2, 4)]
    assert labels == [EDUCATION_GROUP_2_REQUIRED, EDUCATION_GROUP_2_NOT_REQUIRED]


def test_exclusions_sheet_has_sphere_breakdown(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Отсев"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "Сфера" in headers
    assert "Нехватка" in headers
    rows_values = [
        [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)
    ]
    # строка про нехватку в сфере "Право" группы "ВО не требуется"
    pravo_row = next(r for r in rows_values if r[1] == "Право")
    shortfall_col = headers.index("Нехватка")
    assert pravo_row[shortfall_col] == 1


def test_dictionary_sheet_matches_v2_columns(workbook_path):
    from hh_parser.columns_v2 import COLUMNS

    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Словарь колонок"]
    assert ws.max_row - 1 == len(COLUMNS)

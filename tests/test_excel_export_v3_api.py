from openpyxl import load_workbook

from hh_parser.aggregate import SphereCounts
from hh_parser.excel_export_v3_api import export_workbook_v3_api
from hh_parser.rows_v2 import EDUCATION_GROUP_2_NOT_REQUIRED, EDUCATION_GROUP_2_REQUIRED
from hh_parser.rows_v3_api import FinalRowV3Api
from hh_parser.whitelist import SPHERE_IT


def make_row(vacancy_id, employer_id, group, in_subsample) -> FinalRowV3Api:
    return FinalRowV3Api(
        education_group_2=group,
        education_group_filter="higher" if group == EDUCATION_GROUP_2_REQUIRED else "not_required_or_not_specified",
        education_group_filter_label="Высшее" if group == EDUCATION_GROUP_2_REQUIRED else "Не требуется или не указано",
        vacancy_id=vacancy_id,
        link=f"https://hh.ru/vacancy/{vacancy_id}",
        name="Вакансия",
        employer="Компания",
        employer_id=employer_id,
        sector="коммерческий",
        profession_role="Роль",
        role_category=None,
        sphere=SPHERE_IT,
        it_accreditation=False,
        publication_date="2026-10-01",
        salary_from=50000,
        salary_to=70000,
        gross=False,
        salary_net_estimate=60000,
        labor_contract=True,
        civil_law_contract_options=[],
        internship=False,
        work_formats=["ON_SITE"],
        schedule=["FIVE_ON_TWO_OFF"],
        hours=["HOURS_8"],
        dms=False,
        training_mentorship=False,
        education_in_text="не указано",
        education_mismatch=False,
        is_senior_role=False,
        in_matched_subsample=in_subsample,
    )


def test_export_workbook_v3_api_has_all_expected_sheets(tmp_path):
    rows = [
        make_row(1, 10, EDUCATION_GROUP_2_REQUIRED, True),
        make_row(2, 20, EDUCATION_GROUP_2_NOT_REQUIRED, True),
        make_row(3, 30, EDUCATION_GROUP_2_NOT_REQUIRED, False),
    ]
    sphere_rows = [SphereCounts(label=SPHERE_IT, total_found=10, higher=3, special_secondary=2, not_required_or_not_specified=5)]
    self_check_rows = [{"sphere": SPHERE_IT, "api_total": 10, "web_total": 11, "diff_pct": 9.1, "exceeds": False}]
    run_info = {"area": 95, "seed": 42}

    output_path = tmp_path / "test.xlsx"
    export_workbook_v3_api(rows, sphere_rows, {}, {"higher": 1}, {}, {}, self_check_rows, run_info, output_path)

    wb = load_workbook(output_path)
    assert wb.sheetnames == [
        "Выборка",
        "Количества по сферам",
        "Сферы x Группы ВО",
        "Отсев",
        "Самопроверка",
        "Параметры сбора",
        "Словарь колонок",
    ]


def test_export_workbook_v3_api_subsample_column_values_correct(tmp_path):
    rows = [
        make_row(1, 10, EDUCATION_GROUP_2_REQUIRED, True),
        make_row(2, 20, EDUCATION_GROUP_2_NOT_REQUIRED, False),
    ]
    sphere_rows = [SphereCounts(label=SPHERE_IT, total_found=2, higher=1, special_secondary=0, not_required_or_not_specified=1)]
    output_path = tmp_path / "test.xlsx"
    export_workbook_v3_api(rows, sphere_rows, {}, {}, {}, {}, [], {}, output_path)

    wb = load_workbook(output_path)
    ws = wb["Выборка"]
    headers = [cell.value for cell in ws[1]]
    col_idx = headers.index("В подвыборке с одинаковым составом сфер") + 1
    values = [ws.cell(row=r, column=col_idx).value for r in range(2, ws.max_row + 1)]
    assert values == ["да", "нет"]
